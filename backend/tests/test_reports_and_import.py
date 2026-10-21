import io
from datetime import timedelta
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Product, StockMovement, Warehouse
from app.models.base import set_clock, utcnow
from app.services import analytics, inventory
from tests.conftest import fresh


def test_legacy_csv_import(client: TestClient, admin_headers: dict, db: Session) -> None:
    legacy = "Notebook Dell;SN123;TI\nMouse;sn456;Almoxarifado Geral\nlinha quebrada\nRepetido;SN123;TI\n"
    r = client.post(
        "/api/v1/import/legacy-csv",
        headers=admin_headers,
        files={"file": ("estoque.csv", legacy.encode("latin-1"), "text/csv")},
    )
    assert r.status_code == 200, r.text
    result = r.json()
    assert (result["created"], result["skipped"], result["warehouses_created"]) == (2, 1, 2)
    assert result["errors"][0]["line"] == 3

    codes = set(db.scalars(select(Warehouse.code)))
    assert codes == {"TI", "ALMOXARIFADO-GERAL"}
    mouse = db.scalar(select(Product).where(Product.sku == "SN456"))
    assert mouse.total_stock == 1


def test_product_spreadsheet_import_with_brazilian_numbers(
    client: TestClient, admin_headers: dict, db: Session
) -> None:
    csv_text = (
        "sku;nome;categoria;fornecedor;unidade;custo;estoque_minimo;deposito;quantidade\n"
        "A-1;Álcool gel;Limpeza;Fornecedor X;UN;1.234,50;5;Central;10\n"
        ";sem sku;;;;;;;\n"
        "A-2;Luva;EPI;;CX;-3;;;\n"
    )
    r = client.post(
        "/api/v1/import/products",
        headers=admin_headers,
        files={"file": ("produtos.csv", csv_text.encode("utf-8"), "text/csv")},
    ).json()
    assert r["created"] == 1 and r["stock_entries"] == 1
    assert len(r["errors"]) == 2

    product = db.scalar(select(Product).where(Product.sku == "A-1"))
    assert product.cost_price == Decimal("1234.5")
    assert product.category.name == "Limpeza" and product.supplier.name == "Fornecedor X"
    assert product.total_stock == 10

    # reimportar atualiza em vez de duplicar
    update = "sku;nome;estoque_minimo\nA-1;Álcool gel 70%;8\n"
    r2 = client.post(
        "/api/v1/import/products",
        headers=admin_headers,
        files={"file": ("produtos.csv", update.encode("utf-8"), "text/csv")},
    ).json()
    assert r2["updated"] == 1
    fresh(db)
    assert db.scalar(select(Product).where(Product.sku == "A-1")).min_stock == 8


def test_template_roundtrip_xlsx(client: TestClient, admin_headers: dict) -> None:
    template = client.get("/api/v1/import/template?format=xlsx", headers=admin_headers)
    assert template.status_code == 200
    sheet = load_workbook(io.BytesIO(template.content)).active
    assert sheet.cell(row=4, column=1).value == "sku"


def test_exports_have_correct_formats(
    client: TestClient, admin_headers: dict, catalog: dict
) -> None:
    client.post(
        "/api/v1/stock/entries",
        headers=admin_headers,
        json={
            "product_id": catalog["product"].id,
            "warehouse_id": catalog["alm"].id,
            "quantity": 2,
            "unit_cost": 1.5,
        },
    )
    csv_resp = client.get("/api/v1/reports/export/stock-position?format=csv", headers=admin_headers)
    assert csv_resp.content.startswith("﻿".encode())  # BOM para o Excel
    first_line = csv_resp.content.decode("utf-8-sig").splitlines()[0]
    assert first_line.split(";")[0] == "SKU"
    assert "3,00" in csv_resp.content.decode("utf-8-sig")  # valor com vírgula decimal

    for name in ["stock-position", "movements", "abc", "turnover", "replenishment", "low-stock"]:
        pdf = client.get(f"/api/v1/reports/export/{name}?format=pdf", headers=admin_headers)
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF"), name
        xlsx = client.get(f"/api/v1/reports/export/{name}?format=xlsx", headers=admin_headers)
        assert xlsx.status_code == 200 and xlsx.content.startswith(b"PK"), name


def test_labels_pdf(client: TestClient, admin_headers: dict, catalog: dict) -> None:
    r = client.post(
        "/api/v1/labels/pdf",
        headers=admin_headers,
        json={"product_ids": [catalog["product"].id], "copies": 30, "layout": "a4-3x8"},
    )
    assert r.status_code == 200
    assert r.content.count(b"/Type /Page\n") + r.content.count(
        b"/Type /Page\r"
    ) >= 1 or r.content.startswith(b"%PDF")


def test_abc_curve_and_dashboard(db: Session, catalog: dict) -> None:
    alm = catalog["alm"]
    cheap = catalog["product"]
    expensive = Product(sku="CARO", name="Item caro", unit="UN")
    db.add(expensive)
    db.flush()
    for product, cost, qty in [(cheap, "1", 10), (expensive, "100", 10)]:
        inventory.register_entry(
            db,
            product_id=product.id,
            warehouse_id=alm.id,
            quantity=Decimal(20),
            unit_cost=Decimal(cost),
            user=None,
        )
        inventory.register_exit(
            db, product_id=product.id, warehouse_id=alm.id, quantity=Decimal(qty), user=None
        )
    db.commit()

    abc = analytics.abc_curve(db, days=30)
    classes = {i["sku"]: i["abc_class"] for i in abc["items"]}
    assert classes == {"CARO": "A", "P-001": "C"}
    assert abc["total_value"] == 1010.0

    summary = analytics.dashboard_summary(db, days=7)
    assert summary["active_products"] == 2
    assert summary["stock_value"] == 1010.0
    assert summary["series"][-1]["out_qty"] == 20.0
    assert summary["top_exits"][0]["quantity"] == 10.0


def test_movements_outside_period_are_ignored_by_abc(db: Session, catalog: dict) -> None:
    alm, product = catalog["alm"], catalog["product"]
    set_clock(utcnow() - timedelta(days=200))
    try:
        inventory.register_entry(
            db, product_id=product.id, warehouse_id=alm.id, quantity=Decimal(5), user=None
        )
        inventory.register_exit(
            db, product_id=product.id, warehouse_id=alm.id, quantity=Decimal(5), user=None
        )
        db.commit()
    finally:
        set_clock(None)
    assert db.scalar(select(StockMovement.created_at).limit(1)) < utcnow() - timedelta(days=190)
    assert analytics.abc_curve(db, days=90)["total_value"] == 0
