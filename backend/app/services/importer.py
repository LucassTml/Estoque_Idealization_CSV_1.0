"""Importação do CSV da versão 1.0 (C/GTK) e importação/exportação de produtos em CSV ou Excel."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.core.errors import BusinessRuleError
from app.models import Category, Product, Supplier, User, Warehouse
from app.services import inventory

MAX_ROWS = 20_000

# Cabeçalhos aceitos (sem acento, minúsculos) → campo interno
HEADER_ALIASES = {
    "sku": "sku",
    "codigo": "sku",
    "serial": "sku",
    "nome": "name",
    "name": "name",
    "produto": "name",
    "descricao": "description",
    "codigo_barras": "barcode",
    "codigo de barras": "barcode",
    "barcode": "barcode",
    "ean": "barcode",
    "categoria": "category",
    "fornecedor": "supplier",
    "unidade": "unit",
    "un": "unit",
    "custo": "cost_price",
    "preco_custo": "cost_price",
    "preco": "sale_price",
    "preco_venda": "sale_price",
    "estoque_minimo": "min_stock",
    "minimo": "min_stock",
    "estoque_maximo": "max_stock",
    "maximo": "max_stock",
    "deposito": "warehouse",
    "setor": "warehouse",
    "quantidade": "quantity",
    "qtd": "quantity",
}

TEMPLATE_HEADERS = [
    "sku",
    "nome",
    "codigo_barras",
    "categoria",
    "fornecedor",
    "unidade",
    "custo",
    "preco_venda",
    "estoque_minimo",
    "estoque_maximo",
    "deposito",
    "quantidade",
]


class _SemicolonDialect(csv.excel):
    delimiter = ";"


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    warehouses_created: int = 0
    stock_entries: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)


def _strip_accents(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))


def _slug_code(text: str, size: int = 20) -> str:
    code = re.sub(r"[^A-Z0-9]+", "-", _strip_accents(text).upper()).strip("-")
    return (code or "GERAL")[:size]


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BusinessRuleError("Não foi possível ler o arquivo (codificação desconhecida).")


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text:  # formato brasileiro: 1.234,56
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"número inválido: {value!r}") from exc


def _warehouse_for(
    db: Session, name: str, cache: dict[str, Warehouse], result: ImportResult
) -> Warehouse:
    code = _slug_code(name)
    if code in cache:
        return cache[code]
    warehouse = db.scalar(select(Warehouse).where(Warehouse.code == code))
    if warehouse is None:
        warehouse = Warehouse(code=code, name=name.strip()[:80] or code)
        db.add(warehouse)
        db.flush()
        result.warehouses_created += 1
    cache[code] = warehouse
    return warehouse


def import_legacy_csv(db: Session, content: bytes, user: User | None) -> ImportResult:
    """Formato da versão 1.0: `nome;serial;setor`, um item (quantidade 1) por linha."""
    result = ImportResult()
    cache: dict[str, Warehouse] = {}
    lines = _decode(content).splitlines()
    if len(lines) > MAX_ROWS:
        raise BusinessRuleError(f"Arquivo muito grande (máximo {MAX_ROWS} linhas).")
    for number, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        parts = [p.strip() for p in line.split(";")]
        if len(parts) < 3 or not parts[0] or not parts[1]:
            result.errors.append({"line": number, "message": "Esperado: nome;serial;setor"})
            continue
        name, serial, sector = parts[0], parts[1].upper()[:40], parts[2] or "Geral"
        if db.scalar(select(Product.id).where(Product.sku == serial)):
            result.skipped += 1
            continue
        warehouse = _warehouse_for(db, sector, cache, result)
        product = Product(sku=serial, name=name[:160], unit="UN")
        db.add(product)
        db.flush()
        inventory.register_entry(
            db,
            product_id=product.id,
            warehouse_id=warehouse.id,
            quantity=Decimal(1),
            user=user,
            reference="Importação v1.0",
        )
        result.created += 1
        result.stock_entries += 1
    return result


def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip() for h in rows[0]]
        return [
            dict(zip(headers, row, strict=False))
            for row in rows[1:]
            if any(v is not None for v in row)
        ]
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = _SemicolonDialect  # type: ignore[assignment]
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def _normalize_row(raw: dict[str, Any]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for header, value in raw.items():
        if header is None:
            continue
        key = HEADER_ALIASES.get(_strip_accents(str(header)).strip().lower())
        if key:
            row[key] = value.strip() if isinstance(value, str) else value
    return row


def import_products(db: Session, filename: str, content: bytes, user: User | None) -> ImportResult:
    """Cria ou atualiza produtos pelo SKU. Quantidade + depósito geram uma entrada inicial."""
    result = ImportResult()
    raw_rows = _read_rows(filename, content)
    if len(raw_rows) > MAX_ROWS:
        raise BusinessRuleError(f"Arquivo muito grande (máximo {MAX_ROWS} linhas).")
    categories = {c.name.lower(): c for c in db.scalars(select(Category))}
    suppliers = {s.name.lower(): s for s in db.scalars(select(Supplier))}
    warehouses: dict[str, Warehouse] = {}

    for number, raw in enumerate(raw_rows, start=2):
        row = _normalize_row(raw)
        sku = str(row.get("sku") or "").strip().upper()[:40]
        name = str(row.get("name") or "").strip()
        if not sku:
            result.errors.append({"line": number, "message": "SKU obrigatório."})
            continue
        try:
            with db.begin_nested():
                product = db.scalar(select(Product).where(Product.sku == sku))
                is_new = product is None
                if is_new:
                    if not name:
                        raise ValueError("nome obrigatório para produto novo")
                    product = Product(sku=sku, name=name[:160])
                    db.add(product)
                elif name:
                    product.name = name[:160]

                if row.get("barcode"):
                    product.barcode = str(row["barcode"]).strip()[:64]
                if row.get("description"):
                    product.description = str(row["description"])
                if row.get("unit"):
                    product.unit = str(row["unit"]).strip().upper()[:10]
                if row.get("category"):
                    key = str(row["category"]).strip()
                    category = categories.get(key.lower())
                    if category is None:
                        category = Category(name=key[:80])
                        db.add(category)
                        db.flush()
                        categories[key.lower()] = category
                    product.category_id = category.id
                if row.get("supplier"):
                    key = str(row["supplier"]).strip()
                    supplier = suppliers.get(key.lower())
                    if supplier is None:
                        supplier = Supplier(name=key[:120])
                        db.add(supplier)
                        db.flush()
                        suppliers[key.lower()] = supplier
                    product.supplier_id = supplier.id
                for field_name in ("sale_price", "min_stock", "max_stock"):
                    value = _parse_decimal(row.get(field_name))
                    if value is not None:
                        if value < 0:
                            raise ValueError(f"{field_name} não pode ser negativo")
                        setattr(product, field_name, value)
                cost = _parse_decimal(row.get("cost_price"))
                if cost is not None and cost < 0:
                    raise ValueError("custo não pode ser negativo")
                db.flush()

                quantity = _parse_decimal(row.get("quantity"))
                if is_new and quantity and quantity > 0:
                    warehouse = _warehouse_for(
                        db, str(row.get("warehouse") or "Geral"), warehouses, result
                    )
                    inventory.register_entry(
                        db,
                        product_id=product.id,
                        warehouse_id=warehouse.id,
                        quantity=quantity,
                        unit_cost=cost,
                        user=user,
                        reference="Importação de planilha",
                    )
                    result.stock_entries += 1
                elif cost is not None:
                    product.cost_price = cost
                if is_new:
                    result.created += 1
                else:
                    result.updated += 1
        except Exception as exc:  # cada linha é independente graças ao savepoint
            result.errors.append({"line": number, "message": str(exc)[:200]})
    return result


def export_products(db: Session) -> list[dict[str, Any]]:
    products = db.scalars(
        select(Product)
        .options(
            joinedload(Product.category),
            joinedload(Product.supplier),
            joinedload(Product.stock_levels),
        )
        .order_by(Product.sku)
    ).unique()
    return [
        {
            "sku": p.sku,
            "nome": p.name,
            "codigo_barras": p.barcode or "",
            "categoria": p.category.name if p.category else "",
            "fornecedor": p.supplier.name if p.supplier else "",
            "unidade": p.unit,
            "custo": float(p.cost_price),
            "preco_venda": float(p.sale_price),
            "estoque_minimo": float(p.min_stock),
            "estoque_maximo": float(p.max_stock) if p.max_stock is not None else None,
            "estoque_atual": float(p.total_stock),
            "ativo": "sim" if p.is_active else "não",
        }
        for p in products
    ]
