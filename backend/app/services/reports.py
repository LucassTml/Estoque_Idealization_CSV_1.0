"""Exportação de relatórios em Excel (.xlsx), CSV (compatível com o Excel brasileiro) e PDF."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime
from itertools import zip_longest
from typing import Any, Literal

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import PurchaseOrder
from app.services.analytics import local_tz

Format = Literal["xlsx", "csv", "pdf"]
ColumnKind = Literal["text", "qty", "money", "percent", "int", "date"]

MEDIA_TYPES: dict[str, str] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "pdf": "application/pdf",
}

_XLSX_FORMATS = {
    "qty": "#,##0.###",
    "money": '"R$" #,##0.00',
    "percent": "0.0%",
    "int": "0",
}

PO_STATUS_LABELS = {
    "draft": "rascunho",
    "ordered": "enviado ao fornecedor",
    "partially_received": "recebido parcialmente",
    "received": "recebido",
    "cancelled": "cancelado",
}

_LATIN1_REPLACEMENTS = str.maketrans(
    {"—": "-", "–": "-", "→": "->", "•": "-", "“": '"', "”": '"', "’": "'"}
)


@dataclass(frozen=True)
class Column:
    key: str
    title: str
    kind: ColumnKind = "text"
    width: float = 1.0  # peso relativo no PDF


@dataclass
class Report:
    title: str
    columns: list[Column]
    rows: list[dict[str, Any]]
    subtitle: str | None = None
    totals: dict[str, Any] | None = None


def _latin1(text: Any) -> str:
    """As fontes padrão do PDF só cobrem Latin-1 (acentos do português incluídos)."""
    value = "" if text is None else str(text)
    return value.translate(_LATIN1_REPLACEMENTS).encode("latin-1", "replace").decode("latin-1")


def _br_number(value: Any, kind: ColumnKind) -> str:
    if value is None or value == "":
        return ""
    if kind == "money":
        formatted = f"{float(value):,.2f}"
        return "R$ " + formatted.replace(",", "X").replace(".", ",").replace("X", ".")
    if kind == "qty":
        formatted = f"{float(value):,.3f}".rstrip("0").rstrip(".")
        return formatted.replace(",", "X").replace(".", ",").replace("X", ".")
    if kind == "percent":
        return f"{float(value) * 100:.1f}%".replace(".", ",")
    if kind == "int":
        return str(int(value))
    return str(value)


def render(report: Report, fmt: Format) -> bytes:
    if fmt == "xlsx":
        return _to_xlsx(report)
    if fmt == "csv":
        return _to_csv(report)
    return _to_pdf(report)


def _to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([report.subtitle or f"Gerado em {datetime.now(local_tz()):%d/%m/%Y %H:%M}"])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    header_row = 4
    ws.append([c.title for c in report.columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = Alignment(vertical="center")
    for row in report.rows:
        ws.append([row.get(c.key) for c in report.columns])
    if report.totals:
        ws.append([report.totals.get(c.key) for c in report.columns])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    for index, column in enumerate(report.columns, start=1):
        letter = get_column_letter(index)
        number_format = _XLSX_FORMATS.get(column.kind)
        if number_format:
            for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=index, max_col=index):
                cell.number_format = number_format
        values = [str(column.title)] + [str(r.get(column.key) or "") for r in report.rows[:200]]
        ws.column_dimensions[letter].width = min(max(len(v) for v in values) + 3, 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    last = get_column_letter(len(report.columns))
    ws.auto_filter.ref = f"A{header_row}:{last}{header_row + len(report.rows)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _to_csv(report: Report) -> bytes:
    """Separador ";" e BOM UTF-8: o Excel em português abre direto, sem assistente de importação."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([c.title for c in report.columns])
    for row in report.rows:
        writer.writerow(
            [
                _br_number(row.get(c.key), c.kind).removeprefix("R$ ")
                if c.kind in {"money", "qty", "percent"}
                else row.get(c.key, "")
                for c in report.columns
            ]
        )
    return ("﻿" + buffer.getvalue()).encode("utf-8")


class _PDF(FPDF):
    def __init__(self, title: str, subtitle: str | None, orientation: str = "L") -> None:
        super().__init__(orientation=orientation, unit="mm", format="A4")
        self.report_title = _latin1(title)
        self.report_subtitle = _latin1(subtitle or "")
        self.generated = datetime.now(local_tz()).strftime("%d/%m/%Y %H:%M")
        self.set_auto_page_break(auto=True, margin=15)
        self.set_margins(12, 12, 12)

    def header(self) -> None:
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 8, self.report_title, new_x="LMARGIN", new_y="NEXT")
        self.set_font("Helvetica", "", 9)
        self.set_text_color(100)
        self.cell(
            0,
            5,
            self.report_subtitle or f"Gerado em {self.generated}",
            new_x="LMARGIN",
            new_y="NEXT",
        )
        self.set_text_color(0)
        self.ln(3)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(130)
        self.cell(
            0,
            6,
            f"Estoque Idealization  -  {self.generated}  -  página {self.page_no()}/{{nb}}",
            align="C",
        )


def _to_pdf(report: Report) -> bytes:
    pdf = _PDF(report.title, report.subtitle)
    pdf.add_page()
    pdf.set_font("Helvetica", "", 8)
    widths = tuple(c.width for c in report.columns)
    aligns = tuple(
        "RIGHT" if c.kind in {"qty", "money", "percent", "int"} else "LEFT" for c in report.columns
    )
    heading = FontFace(emphasis="BOLD", color=255, fill_color=(30, 58, 138))
    with pdf.table(
        col_widths=widths,
        text_align=aligns,
        headings_style=heading,
        line_height=5,
        cell_fill_color=(243, 244, 246),
        cell_fill_mode="ROWS",
    ) as table:
        table.row([_latin1(c.title) for c in report.columns])
        for row in report.rows:
            table.row([_latin1(_br_number(row.get(c.key), c.kind)) for c in report.columns])
        if report.totals:
            totals_row = table.row()
            for c in report.columns:
                totals_row.cell(
                    _latin1(_br_number(report.totals.get(c.key), c.kind)),
                    style=FontFace(emphasis="BOLD"),
                )
    return bytes(pdf.output())


def purchase_order_pdf(order: PurchaseOrder, company_name: str) -> bytes:
    pdf = _PDF(f"Pedido de compra {order.number}", company_name, orientation="P")
    pdf.add_page()
    supplier = order.supplier
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(95, 6, "Fornecedor")
    pdf.cell(0, 6, "Entrega", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    left = [
        supplier.name,
        f"CNPJ/CPF: {supplier.document or '-'}",
        f"Contato: {supplier.contact_name or '-'}",
        f"E-mail: {supplier.email or '-'}",
        f"Telefone: {supplier.phone or '-'}",
    ]
    right = [
        f"{order.warehouse.code} - {order.warehouse.name}",
        f"Emissão: {(order.ordered_at or order.created_at).astimezone(local_tz()):%d/%m/%Y}",
        f"Previsão: {order.expected_date:%d/%m/%Y}" if order.expected_date else "Previsão: -",
        f"Situação: {PO_STATUS_LABELS.get(order.status.value, order.status.value)}",
    ]
    for a, b in zip_longest(left, right, fillvalue=""):
        pdf.cell(95, 5, _latin1(a)[:60])
        pdf.cell(0, 5, _latin1(b), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Helvetica", "", 9)
    heading = FontFace(emphasis="BOLD", color=255, fill_color=(30, 58, 138))
    with pdf.table(
        col_widths=(1.2, 4, 1.2, 0.8, 1.5, 1.6),
        text_align=("LEFT", "LEFT", "RIGHT", "CENTER", "RIGHT", "RIGHT"),
        headings_style=heading,
        line_height=6,
    ) as table:
        table.row(["SKU", "Produto", "Qtd.", "Un.", "Custo unit.", "Subtotal"])
        for item in order.items:
            table.row(
                [
                    _latin1(item.product.sku),
                    _latin1(item.product.name),
                    _br_number(item.quantity_ordered, "qty"),
                    _latin1(item.product.unit),
                    _br_number(item.unit_cost, "money"),
                    _br_number(item.quantity_ordered * item.unit_cost, "money"),
                ]
            )
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(
        0, 8, f"Total: {_br_number(order.total, 'money')}", align="R", new_x="LMARGIN", new_y="NEXT"
    )
    if order.notes:
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(0, 6, "Observações", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        pdf.multi_cell(0, 5, _latin1(order.notes))
    pdf.ln(18)
    pdf.cell(80, 0, "", border="T")
    pdf.ln(2)
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(80, 4, "Responsável pela compra", align="C")
    return bytes(pdf.output())
